from flask import Flask, render_template, request, redirect, url_for, send_file, flash,session
from werkzeug.security import generate_password_hash,check_password_hash
# <<<<<<< HEAD
from functools import wraps
# =======
# >>>>>>> 516800c66e1ebad7546398b6904d8a63f0a551ed
import sqlite3
import os
import secrets
import uuid
from werkzeug.utils import secure_filename
from reportlab.lib.utils import ImageReader
from io import BytesIO
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfgen import canvas
from flask_wtf.csrf import CSRFProtect
from flask_limiter import Limiter
from flask_limiter.util import get_remote_address
from openpyxl import load_workbook
from datetime import date, datetime

app = Flask(__name__)
<<<<<<< HEAD
# secret key configuration here
csrf = CSRFProtect(app)
limiter = Limiter(
    get_remote_address,
    app=app,
    default_limits=[]
)

@app.errorhandler(429)
def rate_limit_error(error):
    flash(
        "Too many login attempts. Please wait one minute and try again.",
        "error"
    )
    return render_template("login.html"), 429

def login_required(view_function):
    @wraps(view_function)
    def wrapped_view(*args, **kwargs):
        if "admin_id" not in session:
            return redirect(url_for("login"))
        return view_function(*args, **kwargs)

    return wrapped_view


def admin_required(view_function):
    """Allow only administrator accounts to change application data."""
    @wraps(view_function)
    def wrapped_view(*args, **kwargs):
        if "admin_id" not in session:
            return redirect(url_for("login"))
        if session.get("user_role", "admin") != "admin":
            flash("This action is available to administrators only.", "error")
            return redirect(url_for("home"))
        return view_function(*args, **kwargs)
    return wrapped_view
=======
>>>>>>> 516800c66e1ebad7546398b6904d8a63f0a551ed

# Keep the session secret in a local, private file.
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
INSTANCE_DIR = os.path.join(BASE_DIR, "instance")
os.makedirs(INSTANCE_DIR, exist_ok=True)

SECRET_KEY_FILE = os.path.join(INSTANCE_DIR, "secret_key.txt")

if os.path.exists(SECRET_KEY_FILE):
    with open(SECRET_KEY_FILE, "r", encoding="utf-8") as f:
        SECRET_KEY = f.read().strip()
else:
    SECRET_KEY = secrets.token_hex(32)
    with open(SECRET_KEY_FILE, "w", encoding="utf-8") as f:
        f.write(SECRET_KEY)

if not SECRET_KEY:
    raise RuntimeError("Flask secret key is missing.")

app.secret_key = SECRET_KEY
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
<<<<<<< HEAD
app.config["SESSION_COOKIE_SECURE"] = os.environ.get("FLASK_ENV") == "production"
app.config["MAX_CONTENT_LENGTH"] = 3 * 1024 * 1024

DATABASE = os.path.join(BASE_DIR, 'school.db')
UPLOAD_FOLDER = os.path.join(BASE_DIR, "static", "uploads")
=======
app.config["MAX_CONTENT_LENGTH"] = 3 * 1024 * 1024

DATABASE = 'school.db'
UPLOAD_FOLDER = os.path.join("static", "uploads")
>>>>>>> 516800c66e1ebad7546398b6904d8a63f0a551ed
ALLOWED_EXTENSIONS = {"jpg", "jpeg", "png", "webp"}
os.makedirs(UPLOAD_FOLDER, exist_ok=True)


def allowed_file(filename):
    return "." in filename and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS


def save_student_photo(file):
    if not file or not file.filename:
        return None
    if not allowed_file(file.filename):
        flash("Photo not uploaded. Choose a JPG, JPEG, PNG, or WEBP image.", "error")
        return None
    extension = secure_filename(file.filename).rsplit(".", 1)[1].lower()
    filename = uuid.uuid4().hex + "." + extension
    file.save(os.path.join(UPLOAD_FOLDER, filename))
    return "uploads/" + filename


@app.errorhandler(413)
def file_too_large(error):
    flash("That photo is too large. Please choose an image under 3 MB.", "error")
    return redirect(request.referrer or url_for("students_page"))


def connect_database():
    db = sqlite3.connect(DATABASE)
    db.row_factory = sqlite3.Row
    return db


def create_database():
    db = connect_database()
    db.execute('''CREATE TABLE IF NOT EXISTS students (
        id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL,
        father_name TEXT, student_class TEXT, roll_number TEXT,
        date_of_birth TEXT, address TEXT, admission_date TEXT)''')
    columns = [row["name"] for row in db.execute("PRAGMA table_info(students)").fetchall()]
    if "photo" not in columns:
        db.execute("ALTER TABLE students ADD COLUMN photo TEXT")

    if "blood_group" not in columns:
        db.execute("ALTER TABLE students ADD COLUMN blood_group TEXT")

    db.execute('''CREATE TABLE IF NOT EXISTS school (
        id INTEGER PRIMARY KEY CHECK(id=1), name TEXT, address TEXT,
        phone TEXT, email TEXT, principal TEXT, logo TEXT, classes TEXT)''')
    school_columns = [row["name"] for row in db.execute("PRAGMA table_info(school)").fetchall()]
    
    if "logo" not in school_columns:
        db.execute("ALTER TABLE school ADD COLUMN logo TEXT")
    if "classes" not in school_columns:
        db.execute("ALTER TABLE school ADD COLUMN classes TEXT DEFAULT '8,9,10'")
    db.execute("INSERT OR IGNORE INTO school (id, name, address, phone, email, principal, logo, classes) VALUES (1, 'Your School Name', 'School Address', '', '', '', NULL, '8,9,10')")
    db.execute("UPDATE school SET classes='8,9,10' WHERE classes IS NULL OR TRIM(classes)=''")
    db.execute('''
        CREATE TABLE IF NOT EXISTS admins (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
    ''')
    # Existing accounts stay administrators. Add a role column safely for old databases.
    admin_columns = [row["name"] for row in db.execute("PRAGMA table_info(admins)").fetchall()]
    if "role" not in admin_columns:
        db.execute("ALTER TABLE admins ADD COLUMN role TEXT NOT NULL DEFAULT 'admin'")

    # Optional recruiter account. Set both variables in the hosting dashboard;
    # never hard-code the demo password in this source file.
    recruiter_username = os.environ.get("RECRUITER_USERNAME", "").strip()
    recruiter_password = os.environ.get("RECRUITER_PASSWORD", "")
    if recruiter_username and recruiter_password:
        # Create the demo account, or reset its password and role on startup.
        existing = db.execute("SELECT id FROM admins WHERE username=?", (recruiter_username,)).fetchone()
        password_hash = generate_password_hash(recruiter_password)
        if existing:
            db.execute(
                "UPDATE admins SET password_hash=?, role='recruiter' WHERE username=?",
                (password_hash, recruiter_username)
            )
        else:
            db.execute(
                "INSERT INTO admins (username, password_hash, role) VALUES (?, ?, 'recruiter')",
                (recruiter_username, password_hash)
            )
    db.commit()
    db.close()


def get_school():
    db = connect_database()
    school = db.execute('SELECT * FROM school WHERE id=1').fetchone()
    db.close()
    return school


def get_classes():
    school = get_school()
    raw = school['classes'] if school and 'classes' in school.keys() else '8,9,10'
    return [item.strip() for item in (raw or '8,9,10').split(',') if item.strip()]

@app.route("/login", methods=["GET", "POST"])
@limiter.limit("5 per minute")
def login():
    if session.get("admin_id"):
        return redirect(url_for("home"))

    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")

        db = connect_database()
        try:
            admin = db.execute(
                "SELECT * FROM admins WHERE username = ?",
                (username,)
            ).fetchone()
        finally:
            db.close()

        if admin and check_password_hash(
            admin["password_hash"], password
        ):
            session.clear()
            session["admin_id"] = admin["id"]
            session["admin_username"] = admin["username"]
            session["user_role"] = admin["role"] if "role" in admin.keys() else "admin"
            return redirect(url_for("home"))

        flash("Invalid username or password.", "error")

    return render_template("login.html")

@app.route("/logout", methods=["POST"])
@login_required
def logout():
    session.clear()
    flash("You have been logged out successfully.", "success")
    return redirect(url_for("login"))


@app.route('/')
@login_required
def home():
    db = connect_database()
    students = db.execute('SELECT * FROM students ORDER BY id DESC LIMIT 5').fetchall()
    total = db.execute('SELECT COUNT(*) FROM students').fetchone()[0]
    db.close()
    return render_template('index.html', students=students, total_students=total, school=get_school())


@app.route('/students')
@login_required
def students_page():
    db = connect_database()
    students = db.execute('SELECT * FROM students ORDER BY id DESC').fetchall()
    db.close()
    return render_template('students.html', students=students, school=get_school())


@app.route('/id-cards')
@login_required
def id_cards_page():
    db = connect_database()
    students = db.execute('SELECT * FROM students ORDER BY student_class, roll_number, name').fetchall()
    db.close()
    return render_template('id_cards.html', students=students, school=get_school())


@app.route('/transfer-certificates')
@login_required
def transfer_certificates_page():
    db = connect_database()
    students = db.execute('SELECT * FROM students ORDER BY student_class, roll_number, name').fetchall()
    db.close()
    return render_template('transfer_certificates.html', students=students, school=get_school())


@app.route('/add-student', methods=['GET', 'POST'])
@login_required
@admin_required
def add_student():
    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        if name:
            photo = save_student_photo(request.files.get("photo"))
            db = connect_database()
            db.execute('''INSERT INTO students
    (name, father_name, student_class, roll_number, date_of_birth, address, admission_date, photo, blood_group)
    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)''', (
    name, request.form.get('father_name', '').strip(), request.form.get('student_class', ''),
    request.form.get('roll_number', '').strip(), request.form.get('date_of_birth', ''),
    request.form.get('address', '').strip(), request.form.get('admission_date', ''),
    photo, request.form.get('blood_group', '').strip()))
            db.commit()
            db.close()
        return redirect(url_for('students_page'))
    return render_template('add_student.html', school=get_school(), student=None, classes=get_classes())


@app.route('/edit-student/<int:student_id>', methods=['GET', 'POST'])
@login_required
@admin_required
def edit_student(student_id):
    db = connect_database()
    student = db.execute('SELECT * FROM students WHERE id=?', (student_id,)).fetchone()
    if student is None:
        db.close()
        return redirect(url_for('students_page'))
    if request.method == 'POST':
        name = request.form.get('name', '').strip()
        if name:
            new_photo = save_student_photo(request.files.get("photo"))
            photo = new_photo or student["photo"]
            db.execute('''UPDATE students SET name=?, father_name=?, student_class=?, roll_number=?,
    date_of_birth=?, address=?, admission_date=?, photo=?, blood_group=? WHERE id=?''', (
    name, request.form.get('father_name', '').strip(), request.form.get('student_class', ''),
    request.form.get('roll_number', '').strip(), request.form.get('date_of_birth', ''),
    request.form.get('address', '').strip(), request.form.get('admission_date', ''),
    photo, request.form.get('blood_group', '').strip(), student_id))
            db.commit()
        db.close()
        return redirect(url_for('students_page'))
    db.close()
    return render_template('add_student.html', school=get_school(), student=student, classes=get_classes())


@app.route('/delete-student/<int:student_id>', methods=['POST'])
@login_required
@admin_required
def delete_student(student_id):
    db = connect_database()
    student = db.execute("SELECT photo FROM students WHERE id=?", (student_id,)).fetchone()
    if student and student["photo"]:
        photo_path = os.path.join(BASE_DIR, "static", student["photo"])
        if os.path.isfile(photo_path):
            os.remove(photo_path)
    db.execute('DELETE FROM students WHERE id=?', (student_id,))
    db.commit()
    db.close()
    return redirect(url_for('students_page'))


@app.route('/settings', methods=['GET', 'POST'])
@login_required
@admin_required
def settings():
    if request.method == 'POST':
        db = connect_database()
        school = db.execute('SELECT * FROM school WHERE id=1').fetchone()
        logo = school['logo']
        uploaded_logo = request.files.get('logo')
        if uploaded_logo and uploaded_logo.filename:
            if allowed_file(uploaded_logo.filename):
                ext = secure_filename(uploaded_logo.filename).rsplit('.', 1)[1].lower()
                logo_name = 'school_logo_' + uuid.uuid4().hex + '.' + ext
                uploaded_logo.save(os.path.join(UPLOAD_FOLDER, logo_name))
                if logo:
                    old_path = os.path.join(BASE_DIR, 'static', logo)
                    if os.path.isfile(old_path):
                        try: os.remove(old_path)
                        except OSError: pass
                logo = 'uploads/' + logo_name
            else:
                flash('Logo not updated. Choose a JPG, JPEG, PNG, or WEBP image.', 'error')
        classes = request.form.get('classes', '8,9,10').strip() or '8,9,10'
        classes = ','.join(dict.fromkeys(c.strip() for c in classes.split(',') if c.strip())) or '8,9,10'
        db.execute('UPDATE school SET name=?, address=?, phone=?, email=?, principal=?, logo=?, classes=? WHERE id=1', (
            request.form.get('name', '').strip(), request.form.get('address', '').strip(),
            request.form.get('phone', '').strip(), request.form.get('email', '').strip(),
            request.form.get('principal', '').strip(), logo, classes))
        db.commit()
        db.close()
        return redirect(url_for('settings'))
    return render_template('settings.html', school=get_school())


@app.route('/id-card/<int:student_id>')
@login_required
def id_card(student_id):
    db = connect_database()
    student = db.execute('SELECT * FROM students WHERE id=?', (student_id,)).fetchone()
    db.close()
    if student is None:
        return redirect(url_for('students_page'))
    return render_template('id_card.html', student=student, school=get_school())

@app.route('/id-card-pdf/<int:student_id>')
@login_required
def id_card_pdf(student_id):
    db = connect_database()
    student = db.execute(
        'SELECT * FROM students WHERE id=?',
        (student_id,)
    ).fetchone()
    db.close()

    if student is None:
        return redirect(url_for('students_page'))

    school = get_school()
    output = BytesIO()

    # Standard ID-1 card dimensions: 85.6 x 54 mm, converted to points.
    card_width = 242.65
    card_height = 153.07

    pdf = canvas.Canvas(
        output,
        pagesize=(card_width, card_height)
    )

    # White card background.
    pdf.setFillColorRGB(1, 1, 1)
    pdf.rect(0, 0, card_width, card_height, fill=1, stroke=0)

    # Dark school header.
    header_height = 40
    pdf.setFillColorRGB(0.08, 0.16, 0.28)
    pdf.rect(
        0, card_height - header_height,
        card_width, header_height,
        fill=1, stroke=0
    )

    # School logo.
    logo_path = None
    if school['logo']:
        candidate = os.path.join(BASE_DIR, 'static', school['logo'])
        if os.path.isfile(candidate):
            logo_path = candidate

    if logo_path:
        try:
            pdf.drawImage(
                ImageReader(logo_path),
                10, card_height - 32,
                width=22, height=22,
                preserveAspectRatio=True,
                anchor='c',
                mask='auto'
            )
        except Exception:
            logo_path = None

    # School name and address.
    pdf.setFillColorRGB(1, 1, 1)
    pdf.setFont('Helvetica-Bold', 8.5)
    pdf.drawString(
        40 if logo_path else 12,
        card_height - 17,
        (school['name'] or 'Your School Name')[:34]
    )

    pdf.setFont('Helvetica', 6)
    pdf.drawString(
        40 if logo_path else 12,
        card_height - 29,
        (school['address'] or 'School Address')[:43]
    )

    # Card heading.
    pdf.setFillColorRGB(0.19, 0.42, 0.95)
    pdf.setFont('Helvetica-Bold', 6.5)
    pdf.drawString(12, 101, 'OFFICIAL STUDENT ID')

    # Student name.
    pdf.setFillColorRGB(0.10, 0.16, 0.25)
    pdf.setFont('Helvetica-Bold', 10)
    pdf.drawString(
        12, 85,
        (student['name'] or 'Student Name')[:25]
    )

    # Student information.
    pdf.setFillColorRGB(0.28, 0.35, 0.44)
    pdf.setFont('Helvetica', 6.5)

    pdf.drawString(12, 81, 'Student ID: STU-%04d' % student['id'])
    pdf.drawString(12, 69, 'Class: ' + (student['student_class'] or '—')[:20])
    pdf.drawString(12, 57, 'Roll No: ' + (student['roll_number'] or '—')[:20])
    pdf.drawString(12, 45, 'Date of Birth: ' + (student['date_of_birth'] or '—')[:16])
    pdf.drawString(12, 33, 'Blood Group: ' + (student['blood_group'] or '—')[:10])
    pdf.drawString(12, 23, 'Address: ' + (student['address'] or '—')[:27])


    # Student photo.
    photo_x = 187
    photo_y = 36
    photo_width = 43
    photo_height = 58

    photo_path = None
    if student['photo']:
        candidate = os.path.join(BASE_DIR, 'static', student['photo'])
        if os.path.isfile(candidate):
            photo_path = candidate

    if photo_path:
        try:
            pdf.drawImage(
                ImageReader(photo_path),
                photo_x, photo_y,
                width=photo_width,
                height=photo_height,
                preserveAspectRatio=True,
                anchor='c',
                mask='auto'
            )
        except Exception:
            photo_path = None

    if not photo_path:
        pdf.setFillColorRGB(0.91, 0.94, 1)
        pdf.roundRect(
            photo_x, photo_y,
            photo_width, photo_height,
            4, fill=1, stroke=0
        )
        pdf.setFillColorRGB(0.19, 0.42, 0.95)
        pdf.setFont('Helvetica-Bold', 20)
        initial = (student['name'] or 'S')[0].upper()
        pdf.drawCentredString(
            photo_x + photo_width / 2,
            photo_y + photo_height / 2 - 6,
            initial
        )

    # Footer strip.
    pdf.setFillColorRGB(0.94, 0.96, 0.99)
    pdf.rect(1, 1, card_width - 2, 18, fill=1, stroke=0)

    pdf.setFillColorRGB(0.36, 0.42, 0.51)
    pdf.setFont('Helvetica', 6)
    pdf.drawCentredString(
        card_width / 2, 8,
        'Please carry this card on school days.'
    )

    # Thin outer border.
    pdf.setStrokeColorRGB(0.82, 0.86, 0.91)
    pdf.roundRect(
        1, 1,
        card_width - 2, card_height - 2,
        5, fill=0, stroke=1
    )

    pdf.save()
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name='student_id_card_%s.pdf' % student['id'],
        mimetype='application/pdf'
    )

@app.route('/transfer-certificate/<int:student_id>')
@login_required
def transfer_certificate(student_id):
    db = connect_database()
    student = db.execute('SELECT * FROM students WHERE id=?', (student_id,)).fetchone()
    db.close()
    if student is None:
        return redirect(url_for('students_page'))
    return render_template('tc_form.html', student=student, school=get_school())

@app.route('/transfer-certificate-pdf/<int:student_id>', methods=['POST'])
@login_required
def transfer_certificate_pdf(student_id):
    db = connect_database()
    student = db.execute(
        'SELECT * FROM students WHERE id=?', (student_id,)
    ).fetchone()
    db.close()

    if student is None:
        return redirect(url_for('students_page'))

    school = get_school()
    output = BytesIO()
    pdf = canvas.Canvas(output, pagesize=A4)
    width, height = A4

    def value(field, default='Not provided'):
        return request.form.get(field, '').strip() or default

    def draw_label_value(label, text, x, y, label_width=145):
        pdf.setFont('Helvetica-Bold', 9)
        pdf.drawString(x, y, label)
        pdf.setFont('Helvetica', 9)
        pdf.drawString(x + label_width, y, str(text)[:75])

    # Page border
    pdf.setStrokeColorRGB(0.12, 0.22, 0.38)
    pdf.setLineWidth(1.5)
    pdf.rect(32, 32, width - 64, height - 64)

    # School logo
    if school['logo']:
        logo_path = os.path.join(BASE_DIR, 'static', school['logo'])
        if os.path.isfile(logo_path):
            try:
                pdf.drawImage(
                    ImageReader(logo_path),
                    48, height - 112,
                    width=52, height=52,
                    preserveAspectRatio=True,
                    anchor='c',
                    mask='auto'
                )
            except Exception:
                pass

    # School heading
    pdf.setFillColorRGB(0.12, 0.22, 0.38)
    pdf.setFont('Helvetica-Bold', 17)
    pdf.drawCentredString(
        width / 2, height - 76,
        (school['name'] or 'SCHOOL NAME').upper()[:65]
    )

    pdf.setFillColorRGB(0.25, 0.25, 0.25)
    pdf.setFont('Helvetica', 9)
    pdf.drawCentredString(
        width / 2, height - 96,
        (school['address'] or '')[:100]
    )

    pdf.setStrokeColorRGB(0.12, 0.22, 0.38)
    pdf.line(48, height - 120, width - 48, height - 120)

    # Certificate title
    pdf.setFillColorRGB(0.12, 0.22, 0.38)
    pdf.setFont('Helvetica-Bold', 16)
    pdf.drawCentredString(width / 2, height - 150, 'TRANSFER CERTIFICATE')

    pdf.setFillColorRGB(0.15, 0.15, 0.15)
    pdf.setFont('Helvetica', 9)
    pdf.drawString(
        52, height - 180,
        'TC No.: ' + value('tc_number', 'TC-%04d' % student['id'])
    )
    pdf.drawRightString(
        width - 52, height - 180,
        'Date of Issue: ' + value('issue_date')
    )

    pdf.drawString(
        52, height - 198,
        'Admission No.: ' + value('admission_number')
    )
    pdf.drawRightString(
        width - 52, height - 198,
        'Student ID: STU-%04d' % student['id']
    )

    # Student details table
    top = height - 220
    left = 48
    table_width = width - 96
    row_height = 29

    rows = [
        ('1. Student Name', student['name'] or ''),
        ("2. Father's / Guardian's Name", student['father_name'] or ''),
        ('3. Date of Birth', student['date_of_birth'] or ''),
        ('4. Address', student['address'] or ''),
        ('5. Class at Leaving', value('last_class', student['student_class'] or '')),
        ('6. Roll Number', student['roll_number'] or ''),
        ('7. Academic Session', value('academic_session')),
        ('8. Last Date of Attendance', value('last_date')),
        ('9. Subjects Studied', value('subjects')),
        ('10. Result / Promotion Status', value('promotion_status')),
        ('11. Conduct', value('conduct', 'Good')),
        ('12. Reason for Leaving', value('reason')),
        ('13. Fees / Dues Status', value('dues_status')),
        ('14. Date of Application', value('application_date')),
    ]

    for index, (label, text) in enumerate(rows):
        y_top = top - index * row_height
        y_bottom = y_top - row_height

        if index % 2 == 0:
            pdf.setFillColorRGB(0.95, 0.97, 0.99)
            pdf.rect(left, y_bottom, table_width, row_height,
                     fill=1, stroke=0)

        pdf.setStrokeColorRGB(0.78, 0.81, 0.85)
        pdf.rect(left, y_bottom, table_width, row_height,
                 fill=0, stroke=1)
        pdf.line(left + 185, y_bottom, left + 185, y_top)

        pdf.setFillColorRGB(0.12, 0.22, 0.38)
        pdf.setFont('Helvetica-Bold', 8)
        pdf.drawString(left + 7, y_bottom + 10, label[:34])

        pdf.setFillColorRGB(0.15, 0.15, 0.15)
        pdf.setFont('Helvetica', 8)
        pdf.drawString(left + 192, y_bottom + 10, str(text)[:48])

    # Remarks
    remarks_y = top - len(rows) * row_height - 25
    pdf.setFillColorRGB(0.12, 0.22, 0.38)
    pdf.setFont('Helvetica-Bold', 9)
    pdf.drawString(left, remarks_y, 'Remarks:')
    pdf.setFillColorRGB(0.15, 0.15, 0.15)
    pdf.setFont('Helvetica', 9)
    pdf.drawString(left + 60, remarks_y, value('remarks', 'None')[:75])

    # Declaration
    declaration_y = remarks_y - 30
    pdf.setFont('Helvetica', 8)
    pdf.drawString(
        left, declaration_y,
        'Certified that the above particulars have been checked against'
    )
    pdf.drawString(
        left, declaration_y - 12,
        'the school records. Any corrections must be verified before issue.'
    )

    # Signature and seal
    signature_y = 82
    pdf.setStrokeColorRGB(0.45, 0.45, 0.45)
    pdf.line(left + 15, signature_y + 18, left + 145, signature_y + 18)
    pdf.line(width - 205, signature_y + 18, width - 55, signature_y + 18)

    pdf.setFillColorRGB(0.15, 0.15, 0.15)
    pdf.setFont('Helvetica-Bold', 9)
    pdf.drawCentredString(left + 80, signature_y, 'Class Teacher')
    pdf.drawCentredString(width - 130, signature_y, 'Principal')

    pdf.setFont('Helvetica', 8)
    pdf.drawCentredString(width / 2, 55, 'School Seal')

    pdf.save()
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name='transfer_certificate_%s.pdf' % student['id'],
        mimetype='application/pdf'
    )


@app.route('/import-students', methods=['GET', 'POST'])
@login_required
@admin_required
def import_students():
    expected_headers = [
        'name', 'father_name', 'student_class', 'roll_number',
        'date_of_birth', 'address', 'admission_date', 'blood_group'
    ]
    summary = None
    errors = []

    if request.method == 'POST':
        uploaded = request.files.get('excel_file')
        if not uploaded or not uploaded.filename:
            flash('Choose an Excel .xlsx file first.', 'error')
            return redirect(url_for('import_students'))

        if not uploaded.filename.lower().endswith('.xlsx'):
            flash('Please upload an .xlsx Excel file. PDF and .xls files are not supported.', 'error')
            return redirect(url_for('import_students'))

        workbook = None
        try:
            workbook = load_workbook(uploaded, read_only=True, data_only=True)
            sheet = workbook['Students'] if 'Students' in workbook.sheetnames else workbook.active
            rows = sheet.iter_rows(values_only=True)
            header_row = next(rows, None)
            headers = [str(value).strip().lower() if value is not None else '' for value in (header_row or ())]
            if not all(header in headers for header in expected_headers):
                flash('Excel headings do not match the template. Please use the provided Student Import Template.', 'error')
                return redirect(url_for('import_students'))

            indexes = {header: headers.index(header) for header in expected_headers}

            def cell_text(value):
                if value is None:
                    return ''
                if isinstance(value, (datetime, date)):
                    return value.strftime('%d-%m-%Y')
                return str(value).strip()

            db = connect_database()
            imported = 0
            skipped = 0
            valid_classes = set(get_classes())

            try:
                for excel_row_number, values in enumerate(rows, start=2):
                    values = list(values)
                    if not any(value is not None and str(value).strip() for value in values):
                        continue

                    record = {
                        key: cell_text(values[indexes[key]]) if indexes[key] < len(values) else ''
                        for key in expected_headers
                    }
                    name = record['name']
                    student_class = record['student_class']
                    roll_number = record['roll_number']

                    row_errors = []
                    if not name:
                        row_errors.append('student name is required')
                    if not student_class:
                        row_errors.append('class is required')
                    elif student_class not in valid_classes:
                        row_errors.append(
                            'class ' + student_class + ' is not in School Settings (allowed: '
                            + ', '.join(sorted(valid_classes)) + ')'
                        )

                    if row_errors:
                        skipped += 1
                        errors.append('Row ' + str(excel_row_number) + ': ' + '; '.join(row_errors))
                        continue

                    duplicate = None
                    if roll_number:
                        duplicate = db.execute(
                            'SELECT id FROM students WHERE student_class=? AND roll_number=? LIMIT 1',
                            (student_class, roll_number)
                        ).fetchone()

                    if duplicate:
                        skipped += 1
                        errors.append(
                            'Row ' + str(excel_row_number) + ': duplicate roll number '
                            + roll_number + ' in class ' + student_class + '; not imported.'
                        )
                        continue

                    db.execute(
                        "INSERT INTO students "
                        "(name, father_name, student_class, roll_number, date_of_birth, address, admission_date, photo, blood_group) "
                        "VALUES (?, ?, ?, ?, ?, ?, ?, NULL, ?)",
                        (
                            name, record['father_name'], student_class, roll_number,
                            record['date_of_birth'], record['address'], record['admission_date'],
                            record['blood_group']
                        )
                    )
                    imported += 1

                db.commit()
            except Exception:
                db.rollback()
                raise
            finally:
                db.close()

            summary = {'imported': imported, 'skipped': skipped}
        except Exception:
            app.logger.exception('Excel student import failed')
            flash('Could not read this Excel file. Make sure it is a valid .xlsx file using the template.', 'error')
            return redirect(url_for('import_students'))
        finally:
            if workbook is not None:
                workbook.close()

    return render_template(
        'import_students.html',
        school=get_school(),
        summary=summary,
        errors=errors,
        classes=get_classes()
    )


# Initialize/migrate the database for both local runs and Gunicorn deployment.
create_database()

if __name__ == '__main__':
    app.run(debug=False)
